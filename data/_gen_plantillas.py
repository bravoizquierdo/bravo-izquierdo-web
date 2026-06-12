# -*- coding: utf-8 -*-
"""Genera las plantillas Excel (proyectos.xlsx / noticias.xlsx) para OneDrive.
Ejecutar:  python data/_gen_plantillas.py
Se puede borrar; es solo para regenerar las plantillas si hace falta."""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = os.path.dirname(os.path.abspath(__file__))

ACCENT = "D31F26"
DARK = "121317"

hdr_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
hdr_fill = PatternFill("solid", fgColor=DARK)
req_fill = PatternFill("solid", fgColor=ACCENT)
note_font = Font(name="Calibri", italic=True, color="6B7280", size=9)
cell_align = Alignment(vertical="top", wrap_text=True)
thin = Side(style="thin", color="E7E8EC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_sheet(ws, headers, required, notes, widths, rows):
    for c, note in enumerate(notes, start=1):
        cell = ws.cell(row=1, column=c, value=note)
        cell.font = note_font
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.font = hdr_font
        cell.fill = req_fill if h in required else hdr_fill
        cell.alignment = Alignment(vertical="center", horizontal="left")
        cell.border = border
    for r, row in enumerate(rows, start=3):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = cell_align
            cell.border = border
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A3"


# ===================== PROYECTOS =====================
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Proyectos"

p_headers = ["nombre", "categoria", "anio", "ciudad", "descripcion", "descripcion_larga",
             "cliente", "superficie", "estado", "imagen", "imagenes", "enlace", "destacado"]
p_required = {"nombre"}
p_notes = [
    "OBLIGATORIO", "Educacional / Institucional / Hoteles / Civil-Industrial / Inmobiliario",
    "Año entrega", "Ciudad", "Frase corta (tarjeta)", "Texto completo (página de detalle)",
    "Mandante", "m² o magnitud", "Terminado / En ejecución", "URL foto principal",
    "URLs galería separadas por |", "URL externa (opcional)", "si = aparece en portada",
]
p_widths = [34, 18, 8, 16, 40, 52, 18, 14, 16, 30, 40, 30, 10]
p_rows = [
    ["Cuartel SEI Aeródromo Teniente Marsh", "Civil / Industrial", 2020, "Antártica",
     "Infraestructura de extinción de incendios en uno de los entornos más extremos del planeta.",
     "Construcción del cuartel del Servicio de Extinción de Incendios en el Aeródromo Teniente Marsh, en la Antártica chilena. Un proyecto que puso a prueba nuestra capacidad logística y técnica en condiciones climáticas extremas, garantizando seguridad operacional en una de las bases más australes del país.",
     "Fuerza Aérea de Chile", "1.200 m²", "Terminado", "", "", "", "si"],
    ["Facultad de Economía y Empresa UDP", "Educacional", 2013, "Huechuraba",
     "Campus universitario moderno para la Universidad Diego Portales.",
     "Desarrollo del campus de la Facultad de Economía y Empresa de la Universidad Diego Portales, con espacios diseñados para la formación de excelencia: aulas, auditorios, biblioteca y áreas comunes que fomentan la colaboración académica.",
     "Universidad Diego Portales", "8.500 m²", "Terminado", "", "", "", "si"],
    ["Centro Cultural La Moneda", "Institucional", 2006, "Santiago",
     "Uno de los espacios culturales más emblemáticos del país, bajo la Plaza de la Ciudadanía.",
     "Construcción del Centro Cultural La Moneda, ubicado bajo la Plaza de la Ciudadanía frente al Palacio de Gobierno. Una obra de alta complejidad técnica que combina excavación urbana, conservación patrimonial y estándares museográficos internacionales.",
     "DIBAM / Estado de Chile", "7.200 m²", "Terminado", "", "", "", "si"],
    ["Hotel Enjoy Viña del Mar", "Hoteles", 2011, "Viña del Mar",
     "Hotelería de lujo frente a la costa del Pacífico, con los más altos estándares.",
     "Construcción del Hotel Enjoy Viña del Mar, un proyecto de hotelería de lujo frente al océano Pacífico. Terminaciones de primer nivel, áreas de casino, restaurantes y habitaciones con vista al mar, ejecutado bajo exigentes estándares de calidad.",
     "Enjoy S.A.", "32.000 m²", "Terminado", "", "", "", "si"],
    ["Instituto Bicentenario José Miguel Carrera", "Educacional", 2021, "Santiago",
     "El establecimiento educacional más moderno de la región al momento de su entrega.",
     "Construcción del Instituto Bicentenario José Miguel Carrera, considerado el establecimiento educacional más moderno de la región. Infraestructura de vanguardia para la educación pública, con laboratorios, talleres y espacios deportivos de primer nivel.",
     "Ministerio de Educación", "6.800 m²", "Terminado", "", "",
     "https://bravoizquierdo.cl/web/finaliza-construccion-del-nuevo-instituto-bicentenario-jose-miguel-carrera/", "no"],
    ["Obra Connection", "Inmobiliario", 2021, "Santiago",
     "Proyecto distinguido con el Sello Compromiso PRO de la Cámara Chilena de la Construcción.",
     "Desarrollo inmobiliario Connection, distinguido con el Sello Compromiso PRO de la Cámara Chilena de la Construcción por sus buenas prácticas laborales y de seguridad. Un referente en gestión responsable de obra.",
     "Privado", "12.500 m²", "Terminado", "", "",
     "https://bravoizquierdo.cl/web/cchc-presento-sello-compromiso-pro-en-obra-connection/", "no"],
]
style_sheet(ws, p_headers, p_required, p_notes, p_widths, p_rows)

dv_dest = DataValidation(type="list", formula1='"si,no"', allow_blank=True)
ws.add_data_validation(dv_dest)
dv_dest.add("M3:M1000")
dv_est = DataValidation(type="list", formula1='"Terminado,En ejecución"', allow_blank=True)
ws.add_data_validation(dv_est)
dv_est.add("I3:I1000")

wb.save(os.path.join(BASE, "proyectos.xlsx"))
print("proyectos.xlsx OK")

# ===================== NOTICIAS =====================
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Noticias"

n_headers = ["fecha", "titulo", "resumen", "cuerpo", "categoria", "imagen", "enlace"]
n_required = {"fecha", "titulo"}
n_notes = [
    "OBLIGATORIO (dd-mm-aaaa)", "OBLIGATORIO", "Bajada corta (tarjeta)",
    "Texto completo (página de detalle)", "Obras / Calidad / etc.",
    "URL foto", "URL externa (opcional)",
]
n_widths = [16, 42, 46, 60, 18, 30, 30]
n_rows = [
    ["23-09-2021", "Finaliza construcción del nuevo Instituto Bicentenario José Miguel Carrera",
     "El proyecto de Bravo Izquierdo, considerado el establecimiento educacional más moderno de la región, llega a su etapa final.",
     "El nuevo Instituto Bicentenario José Miguel Carrera llega a su etapa final de construcción. Considerado el establecimiento educacional más moderno de la región, el proyecto incorpora laboratorios, talleres especializados y espacios deportivos de primer nivel, reafirmando el compromiso de Bravo Izquierdo con la educación pública de calidad.",
     "Obras", "", "https://bravoizquierdo.cl/web/finaliza-construccion-del-nuevo-instituto-bicentenario-jose-miguel-carrera/"],
    ["24-08-2021", "¿Quieres elegir el mejor lugar para vivir?",
     "Bravo Izquierdo obtiene sello de calidad nuevamente, reafirmando su compromiso con la excelencia.",
     "Bravo Izquierdo obtiene nuevamente el sello de calidad que reconoce la excelencia en sus proyectos inmobiliarios. Este reconocimiento reafirma nuestro compromiso de construir como si fuéramos nuestros propios clientes, cuidando cada detalle de las terminaciones y la experiencia de habitar.",
     "Calidad", "", "https://bravoizquierdo.cl/web/quieres-elegir-el-mejor-lugar-para-vivir/"],
    ["23-06-2021", "Reconocimiento: Premio Mujer Construye 2021",
     "En el marco de la Semana de la Construcción, Bravo Izquierdo participó del Premio Mujer Construye 2021.",
     "En el marco de la Semana de la Construcción 2021, Bravo Izquierdo fue partícipe de la ceremonia del Premio Mujer Construye, iniciativa que releva el rol de las mujeres en la industria de la construcción y promueve la equidad de género en el sector.",
     "Reconocimientos", "", "https://bravoizquierdo.cl/web/reconocimiento-premio-mujer-construye-2021/"],
    ["15-06-2021", "CChC presentó Sello Compromiso PRO en obra Connection",
     "La Cámara Chilena de la Construcción lanzó el Sello Compromiso PRO en nuestra obra Connection.",
     "El 26 de mayo, la Cámara Chilena de la Construcción llevó a cabo el lanzamiento del Sello Compromiso PRO en nuestra obra Connection. Este sello reconoce las buenas prácticas laborales, de seguridad y sostenibilidad en las obras, un ámbito en el que Bravo Izquierdo mantiene un liderazgo constante.",
     "Gremial", "", "https://bravoizquierdo.cl/web/cchc-presento-sello-compromiso-pro-en-obra-connection/"],
    ["07-06-2021", "Bravo Izquierdo participa de la Semana de la Construcción 2021",
     "Bravo Izquierdo participó de la jornada de sostenibilidad de la Semana de la Construcción.",
     "El miércoles 26 de mayo, Bravo Izquierdo participó de la jornada de sostenibilidad de la Semana de la Construcción 2021, compartiendo su experiencia como precursor en sostenibilidad corporativa y gestión de alto desempeño en la industria.",
     "Sostenibilidad", "", "https://bravoizquierdo.cl/web/bravo-izquierdo-participa-de-la-semana-de-la-construccion-2021/"],
    ["20-05-2021", "Certificación ISO 9001:2015 Constructora Bravo Izquierdo",
     "Hace 16 años Bravo Izquierdo certificó su Sistema de Gestión bajo la Norma ISO 9001, hito que renueva año a año.",
     "Hace 16 años, Bravo Izquierdo certificó su Sistema de Gestión a través de la Norma Internacional ISO 9001, hito que renovamos año a año. Esta certificación respalda nuestros procesos de calidad y mejora continua, garantizando estándares de excelencia en cada obra que desarrollamos.",
     "Calidad", "", "https://bravoizquierdo.cl/web/certificacion-iso-90012015-constructora-bravo-izquierdo/"],
]
style_sheet(ws, n_headers, n_required, n_notes, n_widths, n_rows)

wb.save(os.path.join(BASE, "noticias.xlsx"))
print("noticias.xlsx OK")
